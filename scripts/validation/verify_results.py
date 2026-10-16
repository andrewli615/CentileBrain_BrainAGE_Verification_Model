"""Independent saved-result checks, source/archive integrity, and workbook metadata."""

import json
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET

import numpy as np
import pandas as pd
from scipy import stats

from scripts.shared.paths import ROOT, RESULTS, ARCHIVE, PROVENANCE, SOURCE_DIRS
from scripts.shared.data import digest, METHODS, PANELS, age_mask, catalog, training_frames
from scripts.shared.datasets import BRAINAGE_DATASETS


def verify():
    checks=[]

    def same(name,a,b):
        np.testing.assert_allclose(a,b,atol=1e-10,rtol=1e-12,equal_nan=True,err_msg=name)
        checks.append(name)

    inventory=json.loads((ARCHIVE/"migration_inventory.json").read_text())
    for row in inventory["Files"]:
        assert digest(ROOT/row["Archive"])==row["SHA256"],row["Archive"]
    migration_rows=[]
    for row in inventory['Files']:
        active=row['Active']
        if row['Original'] in ('README.md','requirements.txt','.gitignore','plot.py'):
            active=row['Original']
        migration_rows.append(dict(Original=row['Original'],Archive=row['Archive'],Active=active,
                                   Original_SHA256=row['SHA256'],Disposition='Archived and active' if active else 'Archived'))
    pd.DataFrame(migration_rows).to_csv(PROVENANCE/'migration_map.csv',index=False)
    checks.append(f"All {len(inventory['Files'])} historical files preserved byte-for-byte")
    current={str(p.relative_to(ROOT)):digest(p) for name in SOURCE_DIRS for p in (ROOT/name).rglob('*') if p.is_file()}
    protected={k:v for k,v in inventory["Protected_Sources"].items() if Path(k).name!='.DS_Store'}
    substantive={k:v for k,v in current.items() if Path(k).name!='.DS_Store'}
    metadata_changes=[k for k in current.keys() | inventory['Protected_Sources'].keys()
                      if Path(k).name=='.DS_Store' and current.get(k)!=inventory['Protected_Sources'].get(k)]
    assert substantive==protected,"Protected source data files were added, removed, or edited"
    checks.append(f"All {len(substantive)} protected source data files unchanged; Finder metadata tracked separately")
    entries=catalog()
    lookup={(e['Dataset'],e['Sex'],e['Age_Panel'],e['Role']):e for e in entries}
    training=training_frames(entries)
    coefficients=pd.read_csv(RESULTS/"coefficients/training_segments.csv")
    available=pd.read_csv(RESULTS/"summaries/available_statistics.csv")
    matched=pd.read_csv(RESULTS/"summaries/matched_statistics.csv")
    overview=pd.read_csv(RESULTS/"summaries/overview.csv")
    frame_count=0
    for key,label,_ in BRAINAGE_DATASETS:
        for sex,panel in PANELS:
            path=RESULTS/"participants"/key/f"{sex}_{panel}.csv"
            if not path.exists():
                continue
            frame_count+=1
            d=pd.read_csv(path,dtype={'SITE':str,'SubjectID':str})
            assert not d.duplicated(['SITE','SubjectID']).any()
            for role,col in [('Raw','Raw_BrainAGE'),('CentileBrain','CentileBrain_Adjusted')]:
                source=lookup[(key,sex,panel,role)]
                if not source['Path']:
                    continue
                original=pd.read_csv(ROOT/source['Path'],dtype={'SITE':str,'SubjectID':str})
                joined=d.merge(original,on=['SITE','SubjectID'],suffixes=('_export','_source'),validate='one_to_one')
                age=joined.chronological_age_source
                expected=(joined.MR_predicted_age_source-age) if role=='Raw' else joined.Adjusted_BrainAGE
                expected=expected.where(age_mask(age,panel)&np.isfinite(expected))
                same(f'{key} {sex} {panel} source {role}',joined[col],expected)
            for width in (10,5):
                expected=np.full(len(d),np.nan)
                table=coefficients[(coefficients.Sex==sex)&(coefficients.Age_Panel==panel)&(coefficients.Segment_Years==width)]
                ref=training[(sex,panel)]
                for row in table.itertuples():
                    last=row.Bucket_End==(40 if panel=='5-40' else 90)
                    support=age_mask(ref.chronological_age,panel)&np.isfinite(ref.Raw_BrainAGE)&(ref.chronological_age>=row.Bucket_Start)&((ref.chronological_age<=row.Bucket_End) if last else (ref.chronological_age<row.Bucket_End))
                    local=ref[support]
                    same('Training bin count',row.N_Training,len(local))
                    if row.Status!='fitted':
                        continue
                    x=local.chronological_age.to_numpy()
                    slope,intercept=np.linalg.lstsq(np.column_stack([x,np.ones_like(x)]),local.Raw_BrainAGE.to_numpy(),rcond=None)[0]
                    same('Independent training slope',row.Slope_a,slope)
                    same('Independent training intercept',row.Intercept_b,intercept)
                    use=age_mask(d.chronological_age,panel)&np.isfinite(d.Raw_BrainAGE)&(d.chronological_age>=row.Bucket_Start)&((d.chronological_age<=row.Bucket_End) if last else (d.chronological_age<row.Bucket_End))
                    expected[use]=d.Raw_BrainAGE[use]-(slope*d.chronological_age[use]+intercept)
                same(f'{key} {sex} {panel} {width}yr correction',d[f'Training_{width}yr'],expected)
            good=np.isfinite(d.CentileBrain_Adjusted)&age_mask(d.chronological_age,panel)
            if good.sum()>=3 and d.loc[good,'chronological_age'].nunique()>=2:
                ages=d.loc[good,'chronological_age'];values=d.loc[good,'CentileBrain_Adjusted']
                fit=stats.linregress(ages,values)
                same('Independent BrainAGE_S',d.loc[good,'BrainAGE_S'],values-fit.slope*(ages-ages.mean()))
                same('BrainAGE_S preserves mean',d.loc[good,'BrainAGE_S'].mean(),values.mean())
            common=np.isfinite(d[list(METHODS)]).all(axis=1)
            same('Shared participant mask',d.Matched_All_Methods,common)
            for table,use in [(available,np.ones(len(d),dtype=bool)),(matched,common)]:
                for column,method in METHODS.items():
                    rows=d.loc[use&np.isfinite(d[column])]
                    stat=table[(table.Dataset==key)&(table.Sex==sex)&(table.Age_Panel==panel)&(table.Method==method)].iloc[0]
                    same('Whole panel N',stat.N,len(rows))
                    for metric,value in [('MAE',np.mean(np.abs(rows[column]))),('BrainAGE_Mean',rows[column].mean()),('BrainAGE_SD',rows[column].std(ddof=1)),('Age_Mean',rows.chronological_age.mean()),('Age_SD',rows.chronological_age.std(ddof=1)),('Age_Min',rows.chronological_age.min()),('Age_Max',rows.chronological_age.max())]:
                        same(f'Independent {metric}',stat[metric],value)
            over=overview[(overview.Dataset==key)&(overview.Sex==sex)&(overview.Age_Panel==panel)].iloc[0]
            for column,method in METHODS.items():
                same('Overview MAE',over[method],d.loc[common,column].abs().mean())
    # Existing Chinese reference and participant results must remain numerically identical.
    old_root=ARCHIVE/'analyses/training_bucket_bias/outputs'
    old=pd.read_csv(old_root/'training_bucket_coefficients.csv')
    new=coefficients[coefficients.Segment_Years==10]
    for row in old.itertuples():
        found=new[(new.Sex==row.Sex)&(new.Age_Panel==row.Age_Panel)&(new.Bucket_Index==row.Bucket_Index)].iloc[0]
        for metric in ['N_Training','Slope_a','Intercept_b']:
            same('Chinese historical '+metric,found[metric],getattr(row,metric))
    for sex,panel in PANELS:
        old=pd.read_csv(old_root/'self_chinese'/f'{sex}_{panel}'/'participants_training_bucket_adjusted.csv',dtype={'SITE':str,'SubjectID':str})
        new=pd.read_csv(RESULTS/'participants/self_chinese'/f'{sex}_{panel}.csv',dtype={'SITE':str,'SubjectID':str})
        merged=new.merge(old,on=['SITE','SubjectID'],validate='one_to_one')
        same('Chinese historical participant correction',merged.Training_10yr,merged.BrainAGE_TrainingBucketAdjusted)
    verify_plots(checks,training)
    verify_workbook(checks,overview,matched)
    report=dict(Status='passed',Checks_Passed=len(checks),Available_Panels=frame_count,
                Source_Files_Unchanged=len(substantive),Finder_Metadata_Changes=metadata_changes,Historical_Files_Preserved=len(inventory['Files']),
                Notes=['Original inference model is external; exported predictions are the verified inputs.',
                       'Bootstrap/permutation outputs retain original CentileBrain-adjusted inputs.'],Checks=checks)
    (PROVENANCE/'validation_report.json').write_text(json.dumps(report,indent=2)+'\n')
    (PROVENANCE/'migration_report.md').write_text(
        '# Migration and validation report\n\n'
        f'Preserved {len(inventory["Files"])} historical files byte-for-byte and verified all {len(substantive)} protected source data files unchanged.\n\n'
        f'Finder metadata changes (not analysis inputs): {metadata_changes}. These were not restored or edited by the migration.\n\n'
        f'Generated five-method comparisons for {frame_count} available panels. Passed {len(checks)} numerical, source, archive, plot, and workbook checks.\n\n'
        '## Locations\n\n- Sources: original three directories, unchanged.\n- Historical snapshot: archive/pre_reorganization/.\n- Old-to-new map and hashes: archive/pre_reorganization/migration_inventory.json.\n- Workbook and summary tables: results/summaries/.\n- Participant values: results/participants/.\n- Correction coefficients: results/coefficients/.\n- Matched-population figures: results/plots/.\n- Preserved supporting analyses: results/analyses/.\n\n'
        '## Coverage and limitations\n\nSix unavailable sex/age panels are recorded in coverage.csv. Six age-40 Chinese export rows are excluded from older panels. '
        'Unsupported corrections, exclusions, unmatched exports, and training support flags are retained in participant tables and coverage.csv. '
        'The upstream prediction model is external. Corrected MAEs are descriptive gaps, not independent prediction-validation errors. '
        'Historical files may document older paths; active manifests have been rebased. The source catalog is pinned and requires explicit review to change.\n')
    print(f'Verified {len(checks)} checks; {len(substantive)} source data files and {len(inventory["Files"])} archived files unchanged')


def verify_plots(checks,training):
    from scripts.plots import residual_style as style
    import matplotlib.pyplot as plt
    table=pd.read_csv(RESULTS/'summaries/plot_fit_statistics.csv')
    for keys,subset in table.groupby(['Dataset','Sex','Age_Group','Method','Segment_Years'],dropna=False):
        # Saved fit rows are checked against a regression recomputed directly from participant/reference rows.
        key,sex,panel,method,_=keys
        column=next(c for c,label in METHODS.items() if label==method)
        d=pd.read_csv(RESULTS/'participants'/key/f'{sex}_{panel}.csv')
        d=d[d.Matched_All_Methods]
        for row in subset.itertuples():
            ref=training[(sex,panel)] if row.Data=='Training' else d
            value='Raw_BrainAGE' if row.Data=='Training' else column
            ref=ref.loc[age_mask(ref.chronological_age,panel)&np.isfinite(ref[value])]
            if row.Fit_Mode=='independent_segment':
                last=row.Segment_End==(40 if panel=='5-40' else 90)
                ref=ref[(ref.chronological_age>=row.Segment_Start)&((ref.chronological_age<=row.Segment_End) if last else (ref.chronological_age<row.Segment_End))]
            assert row.N==len(ref)
            if len(ref)>=2 and ref.chronological_age.nunique()>=2:
                fit=stats.linregress(ref.chronological_age,ref[value])
                np.testing.assert_allclose([row.Slope,row.Intercept],[fit.slope,fit.intercept],atol=1e-10)
            checks.append('Saved plot fit: '+str(keys))
    # Inspect actual renderer artists for both age panels and every method/segment width.
    for sex,panel in PANELS:
        d=pd.read_csv(RESULTS/'participants/self_chinese'/f'{sex}_{panel}.csv')
        d=d[d.Matched_All_Methods]
        t=training[(sex,panel)]
        t=t.loc[age_mask(t.chronological_age,panel)&np.isfinite(t.Raw_BrainAGE)]
        for column,method in METHODS.items():
            p=style.PlotData('self_chinese','Chinese',sex,panel,
                t[['chronological_age','Raw_BrainAGE']].rename(columns={'chronological_age':'age','Raw_BrainAGE':'residual'}),
                d[['chronological_age',column]].rename(columns={'chronological_age':'age',column:'residual'}))
            for width in (5,10):
                fig,ax=plt.subplots()
                try:
                    style.draw_panel(ax,p,width,show_legend=True)
                    for i,frame in enumerate([p.training,p.observed]):
                        np.testing.assert_allclose(ax.collections[i].get_offsets(),frame[['age','residual']])
                        color=style.TRAINING_SEGMENTED_COLOR if i==0 else style.OBSERVED_SEGMENTED_COLOR
                        lines=[line for line in ax.lines if line.get_color()==color]
                        fits=style.independent_segmented_fit(frame,panel,width)['segments']
                        usable=[s for s in fits if s['can_plot']]
                        assert len(lines)==len(usable)
                        for line,s in zip(lines,usable):
                            np.testing.assert_allclose(line.get_ydata(),s['slope']*line.get_xdata()+s['intercept'])
                        assert all(lines[j].get_xdata()[-1]<lines[j+1].get_xdata()[0] for j in range(len(lines)-1))
                    checks.append(f'Renderer scatter, fits, independent lines: {sex} {panel} {method} {width}')
                finally:
                    plt.close(fig)


def verify_workbook(checks,overview,matched):
    from openpyxl import load_workbook
    path=RESULTS/'summaries/BrainAGE_method_comparison.xlsx'
    wb=load_workbook(path,data_only=True,read_only=False)
    assert wb.sheetnames==['Overview','Matched statistics','Available statistics','Training reference','Methods and coverage']
    for i,row in overview.iterrows():
        for j,method in enumerate(METHODS.values()):
            actual=wb['Overview'].cell(i+6,j+7).value
            expected=row[method]
            if pd.isna(expected):
                assert actual is None
            else:
                np.testing.assert_allclose(actual,expected,atol=1e-10)
            checks.append('Workbook overview matches participants')
    for i,row in matched.iterrows():
        for j,key in enumerate(['Group','Sex','Age_Panel','Method','N','Age_Mean','Age_SD','Age_Min','Age_Max','BrainAGE_Mean','BrainAGE_SD','MAE']):
            actual=wb['Matched statistics'].cell(i+6,j+1).value
            if isinstance(row[key],str):
                assert actual==row[key]
            elif pd.isna(row[key]):
                assert actual is None
            else:
                np.testing.assert_allclose(actual,row[key],atol=1e-10)
    assert wb['Overview'].freeze_panes=='D6'
    assert len(wb['Overview'].tables)==1
    assert next(iter(wb['Overview'].tables.values())).autoFilter is not None
    assert wb['Overview'].page_setup.orientation=='landscape'
    assert wb['Overview'].page_setup.fitToWidth==1 and wb['Overview'].page_setup.fitToHeight==1
    wb.close()
    checks.append('Workbook values, filters, freeze panes, and one-page landscape print settings')


if __name__=='__main__':
    verify()
